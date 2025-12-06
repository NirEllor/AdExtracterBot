import os
import re
from xml.etree.ElementTree import ParseError
import openpyxl
import pandas as pd
from openpyxl.reader.excel import load_workbook

from dropbox_utils import download_excel_from_dropbox, ROOT_IMPORT_PATH
from dropbox.files import WriteMode
from config import dbx


def extract_url_from_formula(formula):
    """Extract URL from Excel HYPERLINK formulas."""
    if not isinstance(formula, str):
        return None
    m = re.search(r'HYPERLINK\("([^"]+)"', formula)
    return m.group(1) if m else None


def filter_failed_files(excel_path, failed_ids_excel, column_name="MASTER CREATIVE ID", attempt=2):
    print(f"\n📝 Filtering Excel (attempt={attempt})...")

    # Load failed creative IDs
    df_failed = pd.read_excel(failed_ids_excel)

    df_failed[column_name] = (
        df_failed[column_name]
        .astype(str)
        .str.replace(r"\.0$", "", regex=True)
        .str.strip()
    )
    failed_ids = set(df_failed[column_name])

    # Choose correct parsing logic:
    if attempt == 2:
        skip_rows = 7        # raw file
        header_row = 8
        data_start = 9
    else:
        skip_rows = 0        # filtered file
        header_row = 1
        data_start = 2

    print(f"skip_rows={skip_rows}, header_row={header_row}, data_start={data_start}")

    # Load via pandas to know which IDs to keep
    df = pd.read_excel(excel_path, skiprows=skip_rows)
    try:
        df[column_name] = df[column_name].astype(str).str.replace(r"\.0$", "", regex=True).str.strip()
        filtered_ids = set(df[df[column_name].isin(failed_ids)][column_name])
        print(f"🔍 Matches found: {len(filtered_ids)}")
        # Load with openpyxl to preserve formulas
        wb = openpyxl.load_workbook(excel_path)
        ws = wb["Report"] if "Report" in wb.sheetnames else wb.active

        new_wb = openpyxl.Workbook()
        new_ws = new_wb.active
        new_ws.title = ws.title

        # Copy header
        for cell in ws[header_row]:
            new_ws.cell(row=1, column=cell.col_idx, value=cell.value)

        # Copy filtered rows
        new_row = 2
        for row in ws.iter_rows(min_row=data_start):
            raw_value = row[1].value
            creative_id = str(raw_value).replace(".0", "").strip() if raw_value else ""

            if creative_id in filtered_ids:
                for cell in row:
                    new_ws.cell(row=new_row, column=cell.col_idx, value=cell.value)
                new_row += 1

        new_wb.save(excel_path)
        print(f"📝 Filtered file saved: {excel_path}")

    except KeyError as e:
        print(e)

    return excel_path

def apply_post_attempt_filtering(report_name, failed_files_excel_name, attempt=2):
    """
    Applies the filtering logic AFTER attempt #1:
    1. Download original REPORT_FILE from Dropbox
    2. Filter it according to FAILED_FILE
    3. Upload filtered version back to Dropbox
    4. Delete temp file
    """
    print("📝 Filtering Excel based on previous failures...")

    # 1 - Download source file from Dropbox
    temp_local_path, _ = download_excel_from_dropbox(
        dbx,
        ROOT_IMPORT_PATH,
        report_name
    )

    if not temp_local_path:
        print("❌ Could not download source report for filtering.")
        return

    # 2 - Apply filtering using FAILED_FILE
    filter_failed_files(
        excel_path=temp_local_path,
        failed_ids_excel=failed_files_excel_name,
        column_name="MASTER CREATIVE ID",
        attempt=attempt
    )

    # 3 - Upload filtered report back to Dropbox
    with open(temp_local_path, "rb") as f:
        dbx.files_upload(
            f.read(),
            f"{ROOT_IMPORT_PATH}/{report_name}",
            mode=WriteMode.overwrite
        )

    print(f"⬆️ Updated filtered report uploaded back to Dropbox.")

    # 4 - Cleanup
    try:
        os.remove(temp_local_path)
        print(f"🗑️ Deleted temporary file: {temp_local_path}")
    except Exception as e:
        print(f"⚠️ Could not remove temp file: {e}")

def extract_external_urls(urls, creative_id_set, excel_path, sheet_name, filtered_excel=False):
    try:
        wb = load_workbook(excel_path)
        ws = wb[sheet_name]

    except Exception as e:
        print("\n❌ ERROR while loading Excel file:", excel_path)

        # בדיקה אם השגיאה קשורה ל-XML
        if isinstance(e.__cause__, ParseError):
            print("XML ParseError:", str(e.__cause__))
        else:
            print("General exception:", str(e))

        # תרצה כאן return כדי לעצור את העיבוד
        return None
    if filtered_excel:
        hidden_rows = {
            row_idx
            for row_idx, row_dim in ws.row_dimensions.items()
            if row_dim.hidden
        }

        visible_count = sum(
            1 for cell in ws['B']
            if cell.row not in hidden_rows
            and cell.row != 1
            and cell.value is not None
        )

        print("number of visible rows (after filter):", visible_count)


    else:
        hidden_rows = None

    pattern = r'=HYPERLINK\("([^"]+)"'


    for url, creative_id in zip(ws['A'], ws['B']):
        if url.row == 1:  # Headers
            continue

        if filtered_excel and url.row in hidden_rows:
            continue

        # 🚫 Case B: explicitly "CREATIVE UNKNOWN"
        raw_value = url.value
        raw_str = str(raw_value).strip()

        if raw_str.lower().startswith("creative unknown"):
            continue

        value = str(url.value)
        match = re.search(pattern, value)

        if match and creative_id.value not in creative_id_set:
            creative_id_set.add(creative_id.value)
            urls[creative_id.value] = match.group(1)
    print("len of creative_id_set is ", len(creative_id_set))
    return None


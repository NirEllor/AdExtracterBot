from openpyxl import load_workbook
import re

def extract(urls, creative_id_set, excel_path, sheet_name, filtered_excel=False):
    wb = load_workbook(excel_path)
    ws = wb[sheet_name]
    if filtered_excel:
        hidden_rows = {
            row_idx
            for row_idx, row_dim in ws.row_dimensions.items()
            if row_dim.hidden
        }

        visible_count = sum(
            1 for cell in ws['B']
            if cell.row not in hidden_rows
            and cell.row != 1
            and cell.value is not None
        )

        print("number of visible rows (after filter):", visible_count)


    else:
        hidden_rows = None

    pattern = r'=HYPERLINK\("([^"]+)"'


    for url, creative_id in zip(ws['A'], ws['B']):
        if url.row == 1:  # Headers
            continue

        if filtered_excel and url.row in hidden_rows:
            continue

        value = str(url.value)
        match = re.search(pattern, value)

        if match and creative_id.value not in creative_id_set:
            creative_id_set.add(creative_id.value)
            urls[creative_id.value] = match.group(1)
    print("len of creative_id_set is ", len(creative_id_set))


if __name__ == '__main__':
    pass
